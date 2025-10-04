from __future__ import annotations

import os
from dataclasses import dataclass
from datetime import datetime
from typing import Optional

from openpyxl import Workbook, load_workbook
from openpyxl.worksheet.worksheet import Worksheet


HEADERS = [
    "Student Name",
    "Day",
    "Accuracy %",
    "WPM",
    "Course Completed Level",
    "Completed Level",
    "OCR Snippet",
    "Image Path",
    "Image Modified",
]


def ensure_workbook(path: str) -> Worksheet:
    os.makedirs(os.path.dirname(path), exist_ok=True)
    if not os.path.exists(path):
        wb = Workbook()
        ws = wb.active
        ws.title = "Progress"
        ws.append(HEADERS)
        wb.save(path)
    wb = load_workbook(path)
    ws = wb.active
    # Ensure headers exist
    if ws.max_row == 0 or [cell.value for cell in ws[1]] != HEADERS:
        if ws.max_row == 0:
            ws.append(HEADERS)
        else:
            for idx, header in enumerate(HEADERS, start=1):
                ws.cell(row=1, column=idx, value=header)
    return ws


def append_row(path: str,
               student_name: str,
               day: Optional[int],
               accuracy: Optional[float],
               wpm: Optional[int],
               course_level: Optional[str],
               completed_level: Optional[str],
               snippet: str,
               image_path: str,
               image_mtime: float) -> None:
    ws = ensure_workbook(path)
    wb = ws.parent

    iso_mtime = datetime.fromtimestamp(image_mtime).isoformat()
    ws.append([
        student_name,
        day,
        accuracy,
        wpm,
        course_level,
        completed_level,
        snippet,
        image_path,
        iso_mtime,
    ])
    wb.save(path)
