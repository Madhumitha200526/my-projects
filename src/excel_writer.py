from datetime import datetime
import os
from typing import Iterable, List
from openpyxl import Workbook, load_workbook
from openpyxl.worksheet.worksheet import Worksheet


_HEADERS = [
    "Student Name",
    "No of Day",
    "Accuracy",
    "Words Per Minute",
    "OCR Snippet",
    "Course Completed Level",
    "Completed Level",
    "Processed At",
    "Image File",
]


class ExcelWriter:
    def __init__(self, spreadsheet_path: str) -> None:
        self.spreadsheet_path = spreadsheet_path
        os.makedirs(os.path.dirname(self.spreadsheet_path), exist_ok=True)
        self._ensure_workbook()

    def _ensure_workbook(self) -> None:
        if not os.path.exists(self.spreadsheet_path):
            wb = Workbook()
            ws = wb.active
            ws.title = "Report"
            ws.append(_HEADERS)
            wb.save(self.spreadsheet_path)
        else:
            # Ensure headers exist (in case of a pre-existing sheet)
            wb = load_workbook(self.spreadsheet_path)
            ws = wb.active
            if ws.max_row == 0:
                ws.append(_HEADERS)
            else:
                # Check if first row matches headers; if not, prepend headers
                first_row = [cell.value for cell in next(ws.iter_rows(min_row=1, max_row=1))]
                if first_row != _HEADERS:
                    ws.insert_rows(1)
                    for idx, header in enumerate(_HEADERS, start=1):
                        ws.cell(row=1, column=idx, value=header)
            wb.save(self.spreadsheet_path)

    def append_rows(self, rows: Iterable[List[str]]) -> int:
        wb = load_workbook(self.spreadsheet_path)
        ws: Worksheet = wb.active
        count = 0
        for row in rows:
            # Ensure the timestamp if not provided in row
            if len(row) < len(_HEADERS):
                now_str = datetime.utcnow().isoformat() + "Z"
                # Extend row to match headers
                extended = list(row) + [None] * (len(_HEADERS) - len(row))
                extended[7] = now_str  # Processed At
                row = extended
            ws.append(row)
            count += 1
        wb.save(self.spreadsheet_path)
        return count
